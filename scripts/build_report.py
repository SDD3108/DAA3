from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
summary = json.loads((ROOT/'results/analysis.json').read_text())
with (ROOT/'results/results.csv').open() as stream:
    rows = list(csv.DictReader(stream))
with (ROOT/'results/bonus.csv').open() as stream:
    bonus = list(csv.DictReader(stream))
assert len(rows) == 50 and len(bonus) == 10

doc = Document()
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
section = doc.sections[0]
section.page_width = Inches(8.27); section.page_height = Inches(11.69)
section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
section.left_margin = Inches(.7); section.right_margin = Inches(.7)
normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.06
for name, size in [('Title',23),('Heading 1',15),('Heading 2',11.5)]:
    style = doc.styles[name]; style.font.name = 'Calibri'; style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_before = Pt(10); style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
footer = section.footer.paragraphs[0]; footer.alignment = 2
run = footer.add_run(); field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); run._r.append(field)
md = []

def heading(text, level=1):
    doc.add_heading(text,level); md.append('#'*(level+1)+' '+text+'\n')

def paragraph(text):
    doc.add_paragraph(text); md.append(text+'\n')

def table(headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        border = OxmlElement('w:'+side); border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'4'); border.set(qn('w:color'),'D9D9D9'); borders.append(border)
    t._tbl.tblPr.append(borders)
    for cell, text in zip(t.rows[0].cells,headers):
        cell.text = text
        for run in cell.paragraphs[0].runs: run.bold = True
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'),'EAF0F5'); cell._tc.get_or_add_tcPr().append(shading)
    for row in data:
        cells = t.add_row().cells
        for cell, text in zip(cells,row): cell.text = str(text)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        for i,cell in enumerate(row.cells):
            if widths: cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(3)
                for run in p.runs: run.font.size = Pt(9)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |')
    md.extend('| '+' | '.join(map(str,row))+' |' for row in data); md.append('')

def image(filename, caption, width=6.75):
    doc.add_picture(str(ROOT/'results/plots'/filename),width=Inches(width))
    p = doc.add_paragraph(caption); p.paragraph_format.space_after = Pt(9)
    for run in p.runs: run.font.size = Pt(9)
    md.append(f'![{caption}](results/plots/{filename})\n')

def page(): doc.add_page_break()

doc.add_paragraph('Graph Algorithms and Performance Analysis',style='Title')
md.append('# Graph Algorithms and Performance Analysis\n')
paragraph('Design and Analysis of Algorithms | Assignment 3')
paragraph('Satimov Damir | SE-2525 | Instructor Taubakabyl Nurlybek')
paragraph('The city network requires minimum-cost connectivity, while the job scheduler requires cycle detection, dependency ordering and paths. The validation suite contains 24 JUnit tests. The submitted benchmark data reproduce all ten reference result triples. Measured traversal counts are linear in the graph size; MST costs also include weight comparisons. Both bonus tasks are implemented.')
heading('Complexity and representation')
paragraph('Let V be vertices and E be input edges. Graph stores adjacency lists, an edge list and maintained in-degrees in O(V + E) space. Undirected input stores two adjacency entries per edge. Extra space below excludes the input graph but includes algorithm work arrays and returned results.')
table(['Operation','Time','Extra space','Reason'],[
['DSU makeSet','O(V)','O(V)','Initialize parent and size arrays.'],
['DSU find and union','O(alpha(V)) amortized','O(1) per call','Size balancing and iterative path compression; DSU storage is O(V).'],
['Kruskal','O(V + E log E)','O(V + E)','Sort edge copy; DSU checks and forest output.'],
['Prim with lazy heap','O(V + E log E)','O(V + E)','At most O(E) queued edges; each removal compares heap keys.'],
['Kosaraju','O(V + E)','O(V + E)','Two iterative DFS passes and a transpose graph.'],
['Kahn','O(V + E)','O(V)','Each vertex is queued once; each outgoing edge decrements one in-degree.'],
['DAG shortest paths','O(V + E)','O(V)','Topological order, then one relaxation pass; distance and parent arrays.'],
['Tarjan bonus','O(V + E)','O(V)','One DFS; index, lowlink, active stack and explicit frames.'],
], [1.35,1.3,1.0,3.1])
heading('What the counters measure',2)
paragraph('dsuOps counts find calls plus successful unions; an attempted union calls find twice. comparisons counts weight comparisons in sorting and the priority queue. edgeScans counts traversal or candidate edges; maxDepth records occupied DFS stack entries, not operations. Prim scans 2E adjacency entries. Kosaraju counts only its two DFS passes (2E), excluding transpose construction. DAG paths report E relaxation scans; their internal Kahn pass is timed but counted separately. Graph construction and validation are outside benchmark timing.')
paragraph('Union by size limits height; path compression flattens later searches. Together they give inverse-Ackermann amortized cost over a sequence of operations, not a strict constant-time guarantee for each individual find.')

page(); heading('Correctness and algorithm selection')
heading('Why the cut property makes MST edges safe',2)
paragraph('Consider a cut that no already selected forest edge crosses. Let e be a lightest edge crossing that cut. If a minimum spanning tree does not contain e, adding e creates a cycle containing another edge f across the same cut. Removing f restores a spanning tree and cannot increase weight because weight(e) is at most weight(f). Thus an optimal tree can contain e together with the existing safe edges. Kruskal applies this argument to DSU components while scanning edges in weight order. Prim applies it to the visited set and its complement, using the smallest crossing edge from its heap. Repeating the argument independently for each connected component proves that both implementations return minimum spanning forests, including graphs with equal or negative weights.')
heading('Why one topological relaxation pass is enough',2)
paragraph('Every finite DAG has a topological order and at least one source in each nonempty dependency region. All sources start at zero, equivalent to adding a virtual super-source with zero-weight outgoing edges. Assume that distances of vertices before v in the order are final. Every incoming edge to v starts earlier, so all candidate final edges of a path to v have already been relaxed. Taking their minimum therefore gives the correct distance to v, establishing the induction step. Processing v propagates that final distance to its successors, and parent pointers record an attaining path. Negative edge weights are safe because a DAG cannot revisit a vertex; in a general graph, reachable negative cycles make shortest walks unbounded, and negative edges also invalidate Dijkstra\'s greedy assumption. The algorithm rejects cycles instead of applying this proof to an unsuitable graph.')
heading('Which algorithm to choose',2)
table(['Situation','Choice and justification'],[
['Sparse undirected graph or existing edge list','Kruskal: sort edges and use DSU; straightforward forest handling.'],
['Undirected adjacency lists already available','Prim: grow from vertex 0, restarting for disconnected components.'],
['Very dense graph','A matrix-based O(V squared) Prim can be preferable; this project uses the required lazy heap, so no universal dense-graph advantage is claimed.'],
['Directed cycles may exist','Kosaraju for clear two-pass SCC decomposition; Tarjan avoids transpose storage and scans edges once.'],
['Dependency order or cycle detection','Kahn: report a cycle when fewer than V vertices leave the queue.'],
['DAG with negative weights','Topological shortest paths: O(V + E), without a heap.'],
['General weighted directed graph','Dijkstra for nonnegative weights; Bellman-Ford if negative edges are possible, with negative-cycle detection.'],
], [2.15,4.6])
paragraph('A deadlock candidate is an SCC with more than one vertex or a singleton with a self-loop. Condensation contracts SCCs into a DAG: a cycle of components would make them mutually reachable and therefore one SCC.')

page(); heading('Measured results and normalized work')
paragraph('Each graph kind starts with a new Random(42). The undirected generator first builds a random spanning tree; the DAG generator orients edges according to shuffled ranks. On an endpoint collision only the second endpoint is redrawn. Each algorithm runs five times, checking that both results and counters repeat. All values below match the assignment reference table.')
reference_rows = []
for i in range(0,len(rows),5):
    group = rows[i:i+5]
    reference_rows.append([f"{int(group[0]['n']):,}",f"{int(group[0]['m']):,}",f"{int(group[0]['result']):,}",group[2]['result'],f"{int(group[4]['result']):,}"])
table(['V','E','MST weight','SCC count','DAG distance sum'],reference_rows)
heading('Counter divided by the complexity formula',2)
ratio_rows = []
for algorithm in ['Kruskal','Prim','Kosaraju','Kahn','DagShortestPath']:
    values=[]
    for density in ['sparse','dense']:
        s=summary[algorithm][density]
        values.append(f"{s['ratio_min']:.3f} to {s['ratio_max']:.3f}")
    ratio_rows.append([algorithm,'comparisons / (E log2 E)' if algorithm in ['Kruskal','Prim'] else 'scans / (V + E)',*values])
table(['Algorithm','Normalization','Sparse range','Dense range'],ratio_rows)
k=summary['Kruskal']
paragraph(f"Kruskal dsuOps/E ranges from {min(k['sparse']['dsu_per_edge']):.3f} to {max(k['sparse']['dsu_per_edge']):.3f} (sparse) and {min(k['dense']['dsu_per_edge']):.3f} to {max(k['dense']['dsu_per_edge']):.3f} (dense). It equals twice the examined-edge fraction plus (V - 1)/E on these connected graphs. These bounded ratios support the expected scaling; finite-size measurements are evidence, not a proof of a complexity class. The comparison ratios need not be exactly constant, particularly with the fixed weight range 1 to 1000.")
heading('Log log slope check',2)
paragraph(f"Least-squares fitting of log(scans) against log(V) gives 1.00 for Prim, Kosaraju, Kahn and DAG relaxation at both densities. Kruskal scan slopes are {k['sparse']['scan_slope']:.2f} (sparse) and {k['dense']['scan_slope']:.2f} (dense); its combined-work slopes are {k['sparse']['operation_slope']:.2f} and {k['dense']['operation_slope']:.2f}. Early termination changes the fraction of edges scanned, whereas E log E describes sorting comparisons. Therefore a scan slope above one must not be attributed directly to sorting. Sorting still considers every input edge.")

page(); heading('Operation growth and scans per edge')
image('operations.png','Figure 1. Sum of dsuOps, edgeScans and comparisons on log-log axes. Kahn and DAG relaxation coincide; Prim and Kosaraju have different operation totals despite equal scan counts.')
paragraph('A separate panel is used for each density. The traversal curves have slope one because their counted scans are exactly proportional to E. MST comparison work is larger than scan work, particularly for the lazy priority queue. The displayed sum measures selected operations, not every CPU instruction or memory access.')
image('scans_per_edge.png','Figure 2. Normalized scans. Prim and Kosaraju overlap at 2; Kahn and DAG relaxation overlap at 1. Kruskal varies because it stops when its forest becomes connected.')
paragraph('Kosaraju uses O(V) explicit stack storage. The dedicated 100,000-vertex path test reaches depth 100,000 without StackOverflowError. Dense random SCCs also produce deep traversals, which is why recursion is avoided.')

page(); heading('MST edge share and practical conclusions')
image('kruskal_share.png','Figure 3. Fraction of sorted edges examined by Kruskal before all vertices are connected.')
paragraph(f"On sparse graphs the examined fraction increases from {k['sparse']['scans_per_edge'][0]:.1%} to {k['sparse']['scans_per_edge'][-1]:.1%}; on dense graphs it ranges from {min(k['dense']['scans_per_edge']):.1%} to {max(k['dense']['scans_per_edge']):.1%}. Kruskal can avoid many DSU checks after obtaining V - 1 edges, but cannot avoid sorting all E edges. Prim always inspects both adjacency entries and manages its frontier in a heap. These data favor Kruskal over this lazy Prim at the largest inputs, without establishing a universal winner for all representations.")
heading('Time measurements',2)
paragraph(f"At V = 100,000 and E = 2,000,000, median times in the recorded run were {float(rows[-5]['time_ms']):.1f} ms for Kruskal, {float(rows[-4]['time_ms']):.1f} ms for Prim, {float(rows[-3]['time_ms']):.1f} ms for Kosaraju, {float(rows[-2]['time_ms']):.1f} ms for Kahn and {float(rows[-1]['time_ms']):.1f} ms for DAG paths. Times on the smallest graphs are noisy and sometimes non-monotonic. JVM compilation, allocation, garbage collection, hardware and background load all affect the clock. Five-run medians reduce outliers but do not remove these effects; correctness values and operation counts are the main evidence. The reproduction commands use a 2 GB maximum heap and require JDK 17 or newer. The benchmark has no separate warm-up phase; machine specifications and the actual Java version were not recorded in the supplied CSV.")
heading('Bonus results and critical path meaning',2)
paragraph('Tarjan returned the same SCC partitions as Kosaraju on all ten benchmark graphs and on 100 random test graphs. It scans E edges instead of 2E: at the largest dense input, 2,000,000 versus 4,000,000 traversal scans. It also avoids building a transpose. Condensation removes internal SCC edges and keeps the maximum parallel transition weight. Longest paths are computed by negating weights as long values and relaxing in topological order; a zero-edge path is allowed. Scheduling uses a longest dependency chain to identify the earliest possible completion bound when edge weights represent nonnegative durations. SCC-internal durations are not modeled. Longest simple paths in unrestricted cyclic graphs are computationally hard; unrestricted positive-cycle walks may have no finite maximum.')
heading('Reproduction and validation',2)
paragraph('Run mvn clean verify, then java -Xmx2g -jar target/daa3-1.0.0.jar and the benchmark.BonusBenchmark class. Run scripts/plot_results.py and scripts/build_report.py to regenerate charts and both reports. README gives complete commands and counter boundaries. Tests include DSU reference labels, 100 random MST comparisons, exhaustive small spanning trees, BFS-based SCC equivalence, 100 Bellman-Ford comparisons with negative DAG weights, path restoration, loops, parallel edges, disconnected graphs and 100,000-vertex DFS paths. See the README for the AI use acknowledgement.')
paragraph('Source: Taubakabyl Nurlybek, Design and Analysis of Algorithms, Assignment 3 - Graphs: DSU, MST, SCC and Shortest Paths, supplied assignment brief. All reported measurements are in results/results.csv; bonus checks are in results/bonus.csv.')

doc.core_properties.title = 'Graph Algorithms and Performance Analysis'
doc.core_properties.subject = 'Design and Analysis of Algorithms Assignment 3'
doc.core_properties.author = 'Satimov Damir'
doc.save(ROOT/'report.docx')
(ROOT/'REPORT.md').write_text('\n'.join(md),encoding='utf-8')
print('Created report.docx and REPORT.md')
